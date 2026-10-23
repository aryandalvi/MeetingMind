"""Text extraction from uploaded TXT / DOCX / PDF (and VTT/SRT caption) files.

Uploads are written to a private temporary directory only for the duration of the extraction
and are always deleted afterwards, including when extraction fails.
"""
from __future__ import annotations

import logging
import os
import re
import shutil
import tempfile
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from ..schemas import ExtractResponse

log = logging.getLogger("meetingmind.files")

TEXT_EXTS = {".txt", ".md", ".vtt", ".srt"}
SUPPORTED_EXTS = TEXT_EXTS | {".docx", ".pdf"}


class FileExtractionError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


@contextmanager
def temporary_upload(data: bytes, suffix: str) -> Iterator[Path]:
    """Write bytes to a temp file inside a private temp dir; remove both on exit."""
    tmp_dir = tempfile.mkdtemp(prefix="meetingmind_")
    path = Path(tmp_dir) / f"upload{suffix}"
    try:
        with open(path, "wb") as fh:
            fh.write(data)
        yield path
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _decode_text(data: bytes) -> str:
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16")
    if b"\x00" in data[:4096]:
        raise FileExtractionError("This file looks like binary data, not a text transcript.", 415)
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _extract_docx(path: Path) -> tuple[str, list[str]]:
    import docx  # python-docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # corrupt or not really a DOCX
        raise FileExtractionError(f"Could not read the Word document ({type(exc).__name__}).", 422) from exc
    lines = [p.text for p in document.paragraphs]
    table_rows = 0
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(": ".join(cells) if len(cells) == 2 else " | ".join(cells))
                table_rows += 1
    warnings = [f"{table_rows} table row(s) were appended after the document text."] if table_rows else []
    return "\n".join(lines), warnings


def _extract_pdf(path: Path) -> tuple[str, list[str], int]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise FileExtractionError("This PDF is password-protected.", 422) from exc
        pages = [page.extract_text() or "" for page in reader.pages]
    except FileExtractionError:
        raise
    except Exception as exc:
        raise FileExtractionError(f"Could not read the PDF ({type(exc).__name__}).", 422) from exc
    text = "\n".join(pages)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join hyphenated line breaks
    warnings = []
    empty_pages = sum(1 for p in pages if not p.strip())
    if empty_pages and empty_pages < len(pages):
        warnings.append(f"{empty_pages} page(s) contained no extractable text (scanned images are not supported).")
    warnings.append("PDF layout can merge or split lines; check speaker labels in the preview.")
    return text, warnings, len(pages)


def extract_text(filename: str, data: bytes, max_bytes: int) -> ExtractResponse:
    name = os.path.basename(filename or "upload")
    ext = Path(name).suffix.lower()
    if ext == ".doc":
        raise FileExtractionError("Legacy .doc files are not supported. Please save the file as .docx.", 415)
    if ext not in SUPPORTED_EXTS:
        raise FileExtractionError(
            f"Unsupported file type '{ext or 'none'}'. Upload a .txt, .docx or .pdf file.", 415
        )
    if not data:
        raise FileExtractionError("The uploaded file is empty.", 400)
    if len(data) > max_bytes:
        raise FileExtractionError(
            f"The file is too large ({len(data) / 1_048_576:.1f} MB; the limit is {max_bytes / 1_048_576:.0f} MB).", 413
        )

    warnings: list[str] = []
    pages = None
    if ext in TEXT_EXTS:
        text = _decode_text(data)
        file_type = "text"
    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise FileExtractionError("This file does not look like a valid .docx document.", 415)
        with temporary_upload(data, ext) as path:
            text, warnings = _extract_docx(path)
        file_type = "docx"
    else:
        if not data.lstrip()[:5].startswith(b"%PDF"):
            raise FileExtractionError("This file does not look like a valid PDF.", 415)
        with temporary_upload(data, ext) as path:
            text, warnings, pages = _extract_pdf(path)
        file_type = "pdf"

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(re.sub(r"\W", "", text)) < 20:
        if file_type == "pdf":
            raise FileExtractionError(
                "No text could be extracted. The PDF is probably scanned (image-based); OCR is not supported.", 422
            )
        raise FileExtractionError("The file contains no readable text.", 422)
    log.info("Extracted %d characters from a %s upload", len(text), file_type)
    return ExtractResponse(
        filename=name,
        file_type=file_type,
        text=text,
        word_count=len(text.split()),
        char_count=len(text),
        pages=pages,
        warnings=warnings,
    )
